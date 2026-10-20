# Append a colour-coded "CATATAN VERIFIKASI" sheet to a filled paper table.
# Copy and adapt: run AFTER writing the metadata values into the template sheet.
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

GREEN = PatternFill("solid", fgColor="C6EFCE")   # valid
AMBER = PatternFill("solid", fgColor="FFEB9C")   # corrected by you
ORANGE = PatternFill("solid", fgColor="FFCC99")  # needs the user's attention


def add_notes(path, rows):
    """rows: list of (paper_label, field, verdict, note)."""
    wb = openpyxl.load_workbook(path)
    if "CATATAN VERIFIKASI" in wb.sheetnames:
        del wb["CATATAN VERIFIKASI"]
    ws = wb.create_sheet("CATATAN VERIFIKASI")
    ws.append(["PAPER", "KOLOM", "STATUS", "CATATAN"])
    for c in range(1, 5):
        ws.cell(1, c).font = Font(bold=True)
    fill = {"valid": GREEN, "diperbaiki": AMBER, "perlu cek": ORANGE}
    for paper, field, verdict, note in rows:
        ws.append([paper, field, verdict, note])
        ws.cell(ws.max_row, 3).fill = fill.get(verdict, ORANGE)
    for col, w in zip("ABCD", (22, 26, 14, 70)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(path)          # ALWAYS save, then re-open to verify
    return path


if __name__ == "__main__":
    demo = [
        ("Paper 1", "HASIL", "valid", "angka cocok dengan PDF"),
        ("Paper 1", "SOFTWARE/APLIKASI", "diperbaiki", "paper pakai Unity DAN Unreal"),
        ("Paper 3", "MASA AKTIF JURNAL", "perlu cek", "periode SK belum terkonfirmasi di ARJUNA"),
    ]
    print(add_notes("/tmp/demo.xlsx", demo))
